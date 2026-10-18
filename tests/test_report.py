"""Sprint Report: the numbers, the Jira call, the dialog and the Excel file."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook

from conftest import SPRINT_REPORTS
from jira_exporter.jira_client import JiraError
from jira_exporter.sprint_report import (
    COLUMNS,
    estimate_field,
    recent_closed_sprints,
    summarize,
    write_report_xlsx,
)
from test_sprint import ROWS, client, write_rows

SPRINT_01 = {"id": 101, "name": "Team Alpha Sprint 01", "state": "closed"}


def test_summarize_counts_committed_completed_spilled_and_adjusted():
    row = summarize("Team Alpha", SPRINT_01, SPRINT_REPORTS[(1, 101)], 1)
    assert (row.team, row.sprint) == ("Team Alpha", "Team Alpha Sprint 01")
    # In the sprint at the start: ABC-1..5, with their estimates then (5+3+5+0+8).
    assert (row.committed, row.committed_points) == (5, 21)
    # Done at the end, added ABC-6 included, with final estimates (8+3+2).
    assert (row.completed, row.completed_points) == (3, 13)
    assert (row.spilled, row.spilled_points) == (2, 5)
    # Added ABC-6, ABC-7; removed ABC-5, ABC-7. ABC-9 (done in another sprint) is ignored.
    assert (row.added, row.removed, row.adjusted) == (2, 2, 4)
    assert row.percent == pytest.approx(13 / 21)
    assert row.values() == ["Team Alpha", "Team Alpha Sprint 01", 5, 21, 3, 13, 2, 5, 4,
                            pytest.approx(13 / 21)]
    assert estimate_field(SPRINT_REPORTS[(1, 101)]) == "customfield_10002"


def test_summarize_empty_report_has_no_percentage():
    row = summarize("T", {"id": 1, "name": "S"}, {"contents": {}})
    assert row.values() == ["T", "S", 0, 0, 0, 0, 0, 0, 0, None]


def test_summarize_keeps_fractional_points():
    report = {"contents": {"completedIssues": [
        {"key": "A-1", "estimateStatistic": {"statFieldValue": {"value": 0.5}}}]}}
    assert summarize("T", {"name": "S"}, report).values()[3] == 0.5


def test_recent_closed_sprints_are_the_boards_own_oldest_first():
    sprints = [
        {"id": 5, "state": "closed", "originBoardId": 1, "completeDate": "2026-03-01"},
        {"id": 3, "state": "closed", "originBoardId": 1, "completeDate": "2026-02-01"},
        {"id": 4, "state": "closed", "originBoardId": 9, "completeDate": "2026-02-15"},
        {"id": 6, "state": "active", "originBoardId": 1},
        {"id": 2, "state": "closed", "completeDate": "2026-01-01"},
    ]
    assert [s["id"] for s in recent_closed_sprints(sprints, 1, 2)] == [3, 5]
    assert [s["id"] for s in recent_closed_sprints(sprints, 1, 10)] == [2, 3, 5]


def test_client_sprint_report(jira_server):
    assert client(jira_server).sprint_report(1, 101) == SPRINT_REPORTS[(1, 101)]
    assert jira_server.requests[-1][0] == \
        "/rest/greenhopper/1.0/rapid/charts/sprintreport?rapidViewId=1&sprintId=101"
    with pytest.raises(JiraError, match="no sprint report for sprint 103 on board 1"):
        client(jira_server).sprint_report(1, 103)


def test_write_report_xlsx(tmp_path):
    rows = [summarize("Team Alpha", SPRINT_01, SPRINT_REPORTS[(1, 101)]),
            summarize("T", {"name": "Empty"}, {"contents": {}})]
    path = tmp_path / "report.xlsx"
    write_report_xlsx(path, rows)
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    assert [c.value for c in ws[2]][:9] == ["Team Alpha", "Team Alpha Sprint 01", 5, 21, 3,
                                            13, 2, 5, 4]
    assert ws.cell(2, 10).value == pytest.approx(13 / 21)
    assert ws.cell(2, 10).number_format == "0%"
    assert ws.cell(3, 10).value is None


# ------------------------------------------------------------------ dialog ---

@pytest.fixture
def dialog(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter.report_dialog import SprintReportDialog

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    settings = QSettings("JiraDesktopExporter", "JiraDesktopExporter")
    settings.setValue("sprint_config", str(write_rows(tmp_path / "teams.xlsx", ROWS)))
    dlg = SprintReportDialog(client(jira_server), settings, None)
    dlg.show()
    yield qapp, dlg
    dlg.reject()


def _run(qapp, dlg):
    from test_ui import wait_until
    dlg.run_btn.click()
    wait_until(qapp, lambda: dlg._worker is None)


def _table(dlg):
    return [[dlg.table.item(r, c).text() for c in range(dlg.table.columnCount())]
            for r in range(dlg.table.rowCount())]


def test_dialog_reports_one_team(dialog):
    qapp, dlg = dialog
    assert [dlg.team.itemText(i) for i in range(dlg.team.count())] == \
        ["All teams", "Team Alpha", "Team Beta", "Team Kanban", "Team Lost"]
    assert not dlg.export_btn.isEnabled()
    dlg.team.setCurrentText("Team Alpha")
    # Team Alpha's ABC row is on board 1; its XYZ row is on board 3, which has no sprints.
    assert dlg.targets() == [("Team Alpha", 1), ("Team Alpha", 3)]
    _run(qapp, dlg)
    assert _table(dlg) == [
        ["Team Alpha", "Team Alpha Sprint 01", "5", "21", "3", "13", "2", "5", "4", "62%"],
        ["Team Alpha", "Team Alpha Sprint 02", "2", "8", "2", "8", "0", "0", "0", "100%"],
    ]
    status = dlg.status.text()
    assert "2 sprints reported" in status and "(Story Points)" in status
    assert "Team Alpha, board 3: no closed sprints." in status
    assert dlg.export_btn.isEnabled()
    assert dlg.settings.value("report_team") == "Team Alpha"


def test_dialog_all_teams_lists_boards_it_could_not_read(dialog):
    qapp, dlg = dialog
    dlg.team.setCurrentText("All teams")
    dlg.count.setValue(1)
    _run(qapp, dlg)
    assert [r[1] for r in _table(dlg)] == ["Team Alpha Sprint 02"]
    status = dlg.status.text()
    assert "Team Beta, board 2: no closed sprints." in status
    assert "Team Kanban, board 7:" in status and "Team Lost, board 99:" in status


def test_dialog_exports_the_table(dialog, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    qapp, dlg = dialog
    dlg.team.setCurrentText("Team Alpha")
    _run(qapp, dlg)
    out = tmp_path / "out"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    dlg.export_btn.click()
    ws = load_workbook(tmp_path / "out.xlsx").active
    assert ws.max_row == 3 and ws.cell(3, 2).value == "Team Alpha Sprint 02"


def test_dialog_sorts_percentage_as_a_number(dialog):
    from PySide6.QtCore import Qt
    qapp, dlg = dialog
    dlg.team.setCurrentText("Team Alpha")
    _run(qapp, dlg)
    dlg.table.sortItems(len(COLUMNS) - 1, Qt.DescendingOrder)
    assert [r[-1] for r in _table(dlg)] == ["100%", "62%"]


def test_dialog_missing_config(dialog, tmp_path):
    qapp, dlg = dialog
    dlg.config_path.setText(str(tmp_path / "nope.xlsx"))
    dlg.load_config()
    assert "Sprint config not found" in dlg.status.text()
    assert not dlg.run_btn.isEnabled()


def test_menu_opens_sprint_report_only_when_logged_in(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter import app as appmod
    from test_ui import wait_until

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    assert not win.sprint_report_action.isEnabled()
    win.sprint_report_action.trigger()
    assert win.report_dialog is None

    win.login_action.trigger()
    win.login.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    win.login.username.setText("alice")
    win.login.password.setText("secret")
    win.login.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    win.sprint_report_action.trigger()
    dlg = win.report_dialog
    assert dlg is not None and dlg.isVisible() and dlg.client is win.export.client
    win.sprint_report_action.trigger()
    assert win.report_dialog is dlg

    win.logout_action.trigger()
    wait_until(qapp, lambda: win.report_dialog is None)
    assert not win.sprint_report_action.isEnabled()
    win.close()
