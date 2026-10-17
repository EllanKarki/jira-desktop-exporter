"""Drives the real window headlessly: log in, export, check the file."""
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import load_workbook
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication, QMessageBox

from jira_exporter import app as appmod

ROOT = Path(__file__).resolve().parents[1]


def wait_until(qapp, cond, timeout=10):
    import time
    end = time.time() + timeout
    while not cond() and time.time() < end:
        qapp.processEvents()
        time.sleep(0.02)
    assert cond()


def test_login_then_export(jira_server, tmp_path, monkeypatch):
    # Keep the test's saved settings out of the real user profile.
    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()

    win.login_action.trigger()
    lp = win.login
    lp.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    lp.username.setText("alice")
    lp.password.setText("wrong")
    lp.connect_btn.click()
    wait_until(qapp, lambda: "Login failed" in lp.status.text())

    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    assert win.stack.currentWidget() is win.start
    win.jql_action.trigger()
    assert lp.password.text() == ""

    ep = win.export
    ep.jql.setPlainText("project = ABC")
    ep.config_path.setText(str(ROOT / "config" / "fields.json"))
    out = tmp_path / "ui.xlsx"
    ep.out_path.setText(str(out))
    ep.export_btn.click()
    wait_until(qapp, lambda: "Done: 250 issues" in ep.log.toPlainText())
    log = ep.log.toPlainText()
    assert "Warning: Field name 'Priority' was not found" in log
    ws = load_workbook(out).active
    assert ws.max_row == 251
    assert ws["A1"].value == "Issue Key"


def test_cloud_login_then_export(cloud_server, tmp_path, monkeypatch):
    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()

    win.login_action.trigger()
    lp = win.login
    lp.url.setText("https://acme.atlassian.net")
    lp._guess_deployment()
    assert lp.deployment.currentData() == "cloud"
    assert lp.user_label.text() == "Email"
    assert lp.secret_label.text() == "API token"
    assert lp.use_token.isHidden()

    lp.url.setText(f"http://127.0.0.1:{cloud_server.server_address[1]}")  # stays on Cloud
    lp.username.setText("alice")
    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    assert win.stack.currentWidget() is win.start
    win.jql_action.trigger()
    assert "Jira Cloud" in win.export.user_label.text()

    ep = win.export
    ep.jql.setPlainText("project = ABC")
    ep.config_path.setText(str(ROOT / "config" / "fields.json"))
    out = tmp_path / "cloud_ui.xlsx"
    ep.out_path.setText(str(out))
    ep.export_btn.click()
    wait_until(qapp, lambda: "Done: 250 issues" in ep.log.toPlainText())
    assert load_workbook(out).active.max_row == 251


def _fresh_settings(tmp_path):
    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    return QSettings("JiraDesktopExporter", "JiraDesktopExporter")


def test_type_guessed_from_settings_saved_before_cloud_support(tmp_path):
    s = _fresh_settings(tmp_path)
    s.setValue("url", "https://acme.atlassian.net")
    s.setValue("use_token", True)
    s.sync()
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.login_action.trigger()
    lp = win.login
    assert lp.deployment.currentData() == "cloud"
    assert lp.secret_label.text() == "API token"


def test_saved_type_wins_over_url_guess(tmp_path):
    s = _fresh_settings(tmp_path)
    s.setValue("url", "https://jira.acme.com")
    s.setValue("deployment", "cloud")  # Cloud on a custom domain, chosen by hand
    s.sync()
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.login_action.trigger()
    lp = win.login
    assert lp.deployment.currentData() == "cloud"
    lp.url.setText("https://jira2.acme.com")  # editing the URL keeps the manual choice
    lp._guess_deployment()
    assert lp.deployment.currentData() == "cloud"


def test_url_guess_switches_back_only_when_it_chose_cloud(tmp_path):
    _fresh_settings(tmp_path)
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.login_action.trigger()
    lp = win.login
    lp.url.setText("https://acme.atlassian.net")
    lp._guess_deployment()
    assert lp.deployment.currentData() == "cloud"
    lp.url.setText("https://jira.acme.com")
    lp._guess_deployment()
    assert lp.deployment.currentData() == "datacenter"
    # Picking Cloud by hand for a custom domain sticks through URL edits.
    lp.deployment.setCurrentIndex(lp.deployment.findData("cloud"))
    lp.deployment.activated.emit(lp.deployment.currentIndex())
    lp.url.setText("https://jira.acme.org")
    lp._guess_deployment()
    assert lp.deployment.currentData() == "cloud"


def test_progress_without_total_stays_busy(tmp_path):
    _fresh_settings(tmp_path)
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    ep = win.export
    ep._on_progress(100, 0)
    assert ep.progress.maximum() == 0
    assert ep.progress_label.text() == "100 issues fetched (total unknown)"
    ep._on_progress(100, 250)
    assert ep.progress.maximum() == 250
    assert ep.progress_label.text() == "100 of 250 issues fetched"


def test_guess_deployment_survives_malformed_saved_url(tmp_path):
    s = _fresh_settings(tmp_path)
    s.setValue("url", "https://jira.acme.com]")
    s.setValue("deployment", "cloud")
    s.sync()
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    assert win.login.deployment.currentData() == "cloud"


def test_url_round_trip_keeps_manual_cloud_choice(tmp_path):
    _fresh_settings(tmp_path)
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.login_action.trigger()
    lp = win.login
    lp.url.setText("https://jira.acme.com")
    lp._guess_deployment()
    lp.deployment.setCurrentIndex(lp.deployment.findData("cloud"))
    lp.deployment.activated.emit(lp.deployment.currentIndex())
    for url in ("https://acme.atlassian.net", "https://jira.acme.com"):
        lp.url.setText(url)
        lp._guess_deployment()
        assert lp.deployment.currentData() == "cloud"


def test_cancelled_export_does_not_show_full_bar(tmp_path, monkeypatch):
    _fresh_settings(tmp_path)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    ep = win.export

    class FakeWorker:
        stopped_early = True

    class FakeThread:
        def quit(self):
            pass

    ep._worker, ep._thread = FakeWorker(), FakeThread()
    ep._on_progress(100, 250)
    ep._on_finished(100, "/tmp/x.xlsx")
    assert ep.progress.value() == 100 and ep.progress.maximum() == 250
    assert ep.progress_label.text() == "Cancelled: 100 issues saved"


def test_cancel_during_last_page_reports_complete_export(jira_server, tmp_path, monkeypatch):
    """Cancel only counts when pages were actually skipped."""
    from jira_exporter.config import load_config
    from jira_exporter.jira_client import Credentials, JiraClient

    client = JiraClient(Credentials(f"http://127.0.0.1:{jira_server.server_address[1]}",
                                    "alice", "secret"))
    cfg = load_config(ROOT / "config" / "fields.json")
    worker = appmod.ExportWorker(client, cfg, "project = ABC", tmp_path / "o.xlsx")
    worker.progress.connect(lambda done, total: done >= 250 and worker.cancel())
    done = []
    worker.finished.connect(lambda n, p: done.append(n))
    worker.run()
    assert done == [250] and worker._cancel and not worker.stopped_early

    worker = appmod.ExportWorker(client, cfg, "project = ABC", tmp_path / "p.xlsx")
    worker.progress.connect(lambda done, total: done >= 100 and worker.cancel())
    done = []
    worker.finished.connect(lambda n, p: done.append(n))
    worker.run()
    assert done == [100] and worker.stopped_early


def test_relogin_clears_previous_result(tmp_path):
    _fresh_settings(tmp_path)
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    ep = win.export
    ep._on_progress(250, 250)
    ep.progress_label.setText("Done: 250 issues saved")

    class FakeClient:
        base_url, is_cloud = "https://jira.acme.com", False

    ep.set_session(FakeClient(), {"displayName": "Bob"})
    assert ep.progress_label.text() == ""
    assert ep.progress.value() == 0


def test_menus_follow_login_state(jira_server, tmp_path, monkeypatch):
    _fresh_settings(tmp_path)
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    titles = [a.text().replace("&", "") for a in win.menuBar().actions()]
    assert titles == ["Jira", "Configuration", "Extract", "Jira Actions"]
    jira_items = [a.text().replace("&", "") for a in win.menuBar().actions()[0].menu().actions()
                  if not a.isSeparator()]
    assert jira_items == ["Login", "Logout", "Exit"]

    # Logged out: only Login and Exit are available.
    assert win.login_action.isEnabled() and win.exit_action.isEnabled()
    for a in (win.logout_action, win.discover_action, win.jql_action):
        assert not a.isEnabled()

    win.login_action.trigger()
    lp = win.login
    lp.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    lp.username.setText("alice")
    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    assert win.stack.currentWidget() is win.start
    win.jql_action.trigger()
    assert not win.login_action.isEnabled()
    for a in (win.logout_action, win.discover_action, win.jql_action):
        assert a.isEnabled()

    win.discover_action.trigger()
    dlg = win.export.fields_dialog
    assert dlg is not None and dlg.isVisible()
    dlg.reject()

    # Logout is greyed out while an export runs.
    win.export.export_btn.setEnabled(False)
    win._update_menus()
    assert not win.logout_action.isEnabled()
    win.export.export_btn.setEnabled(True)

    win.jql_action.trigger()
    assert win.stack.currentWidget() is win.export
    win.logout_action.setEnabled(True)
    win.logout_action.trigger()
    assert win.stack.currentWidget() is win.start and win.export.client is None
    assert win.login_action.isEnabled() and not win.jql_action.isEnabled()
    win.login_action.trigger()
    assert win.stack.currentWidget() is win.login

    win.show()
    win.exit_action.trigger()
    assert not win.isVisible()


def test_login_form_only_opens_from_the_menu(tmp_path):
    s = _fresh_settings(tmp_path)
    s.setValue("url", "https://jira.acme.com")
    s.setValue("username", "alice")
    s.sync()
    QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.show()
    assert win.stack.currentWidget() is win.start
    assert not win.login.isVisible()
    assert "Jira &gt; Login" in win.start.findChildren(appmod.QLabel)[1].text()
    win.login_action.trigger()
    assert win.stack.currentWidget() is win.login and win.login.isVisible()
    assert win.login.password.hasFocus() or not win.isActiveWindow()
    win.close()


def test_export_form_only_opens_from_the_menu(jira_server, tmp_path):
    _fresh_settings(tmp_path)
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    assert not hasattr(win.export, "discover_btn")
    win.login_action.trigger()
    lp = win.login
    lp.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    lp.username.setText("alice")
    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    assert win.stack.currentWidget() is win.start
    assert "Connected to" in win.start.hint.text() and "Alice A" in win.start.hint.text()
    win.jql_action.trigger()
    assert win.stack.currentWidget() is win.export
    win.logout_action.trigger()
    assert win.stack.currentWidget() is win.start
    assert "not connected" in win.start.hint.text()
