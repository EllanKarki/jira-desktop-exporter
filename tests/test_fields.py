"""'Discover fields': scanning a project's fields and saving a fields.json from them."""
import json
import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook

from conftest import ADF_DESC, FIELDS
from jira_exporter.config import load_config, resolve_fields
from jira_exporter.field_discovery import (
    FieldUsage,
    config_entries,
    discover_fields,
    has_value,
    jql_for_project,
    write_config,
    write_field_list_xlsx,
)
from jira_exporter.jira_client import CLOUD, Credentials, JiraClient, JiraError

ROOT = Path(__file__).resolve().parents[1]


def client(srv, deployment="datacenter"):
    return JiraClient(Credentials(f"http://127.0.0.1:{srv.server_address[1]}", "alice", "secret",
                                  deployment=deployment))


@pytest.mark.parametrize("value, expected", [
    (None, False), ("", False), ("  ", False), ("x", True), (0, True), (False, True),
    ([], False), (["a"], True), ({}, False), ({"name": "Open"}, True),
    ({"type": "doc", "version": 1, "content": []}, False), (ADF_DESC, True),
    ({"comments": [], "maxResults": 0, "total": 0, "startAt": 0}, False),
    ({"comments": [{"body": "hi"}], "total": 1}, True),
    ({"self": "x", "watchCount": 0, "isWatching": False}, False),
    ({"self": "x", "watchCount": 2, "isWatching": False}, True),
    ({"self": "x", "votes": 0, "hasVoted": False}, False),
    ("{}", False),  # Cloud's Development field when empty
    ({"issuerestrictions": {}, "shouldDisplay": False}, False),
    ({"showField": False}, False),
    ({"self": "x", "value": "High", "id": "1"}, True),
    ({"outer": {"inner": None}}, False),
])
def test_has_value(value, expected):
    assert has_value(value) is expected


def test_jql_quotes_the_key():
    assert jql_for_project("ABC") == 'project = "ABC" ORDER BY updated DESC'
    assert jql_for_project('A"B') == 'project = "A\\"B" ORDER BY updated DESC'


def _by_id(result):
    return {f.id: f for f in result.fields}


@pytest.mark.parametrize("cloud", [False, True])
def test_discover_fields(jira_server, cloud_server, cloud):
    srv = cloud_server if cloud else jira_server
    c = client(srv, CLOUD if cloud else "datacenter")
    steps = []
    result = discover_fields(c, " ABC ", sample_size=120, progress=steps.append)

    assert (result.project_key, result.project_name, result.sampled) == ("ABC", "Alpha Beta", 120)
    assert result.warnings == []
    f = _by_id(result)
    assert f["issuekey"].used_in == 120 and result.fields[0].is_key
    assert f["summary"].used_in == 120
    assert f["summary"].issue_types == {"Bug", "Story"}
    assert f["assignee"].used_in == 80  # every third issue is unassigned
    assert f["description"].used_in == 120  # rich text on Cloud, plain on DC
    assert f["customfield_10003"].used_in == 0 and f["customfield_10003"].issue_types == {"Bug"}
    assert f["customfield_10003"].type == "option" and f["customfield_10003"].custom
    relevant = {x.id for x in result.relevant}
    assert "customfield_10003" in relevant  # offered on a create screen, never filled
    assert "customfield_10099" not in relevant  # neither
    assert f["labels"].type == "array"
    # The scan pages through create metadata (2 per page in the fake) and asks for all fields.
    paths = [p for p, _ in srv.requests]
    assert sum("/issuetypes/1?" in p for p in paths) == 2
    searches = [b for p, b in srv.requests if b and "jql" in b and "fields" in b]
    assert searches and all(b["fields"] == ["*all"] for b in searches)
    assert any("Checked 100 issues" in s for s in steps)


def test_discover_without_sampling(jira_server):
    result = discover_fields(client(jira_server), "ABC", sample_size=0)
    assert result.sampled == 0
    assert not any(b for _, b in jira_server.requests if b and "jql" in b)
    assert {x.id for x in result.relevant} == {
        "issuekey", "summary", "description", "duedate", "customfield_10003", "labels",
        "customfield_10002"}


def test_discover_when_createmeta_is_missing(jira_server):
    jira_server.createmeta_status = 404  # Data Center before 8.4
    result = discover_fields(client(jira_server), "ABC", sample_size=50)
    assert result.sampled == 50
    assert "Data Center 8.4" in result.warnings[0]
    assert not any(f.issue_types for f in result.fields)
    assert _by_id(result)["summary"].used_in == 50


def test_discover_when_createmeta_is_forbidden(jira_server):
    jira_server.createmeta_status = 403
    result = discover_fields(client(jira_server), "ABC", sample_size=10)
    assert "Access denied by Jira (403)" in result.warnings[0]
    assert "\n" not in result.warnings[0]


def test_discover_unknown_project(jira_server):
    with pytest.raises(JiraError, match="Project 'ZZZ' was not found"):
        discover_fields(client(jira_server), "ZZZ")


def test_discover_can_be_cancelled(jira_server):
    seen = []
    result = discover_fields(client(jira_server), "ABC", sample_size=250,
                             progress=seen.append,
                             cancelled=lambda: any("Checked 100" in s for s in seen))
    assert result.sampled == 100


def test_config_round_trip(jira_server, tmp_path):
    result = discover_fields(client(jira_server), "ABC", sample_size=30)
    chosen = [f for f in result.fields if f.id in ("issuekey", "summary", "customfield_10002")]
    template = tmp_path / "old.json"
    template.write_text(json.dumps({"max_results_per_page": 50, "fields": [{"name": "X"}],
                                    "note": "kept"}), encoding="utf-8")
    out = tmp_path / "new.json"
    write_config(out, chosen, template=template)

    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["note"] == "kept" and data["max_results_per_page"] == 50
    assert data["fields"] == [{"id": "key", "header": "Issue Key"},
                              {"id": "customfield_10002", "header": "Story Points"},
                              {"id": "summary", "header": "Summary"}]
    resolved, problems = resolve_fields(load_config(out).fields, FIELDS)
    assert problems == [] and [r.header for r in resolved] == [
        "Issue Key", "Story Points", "Summary"]


def test_write_config_without_template(tmp_path):
    out = tmp_path / "f.json"
    write_config(out, [FieldUsage("customfield_1", "Équipe")], template=tmp_path / "missing.json")
    text = out.read_text(encoding="utf-8")
    assert "Équipe" in text  # names stay readable, not \u-escaped
    assert load_config(out).page_size == 100


def test_config_entries_key_alias():
    assert config_entries([FieldUsage("key", "Key")]) == [{"id": "key", "header": "Issue Key"}]


def test_field_list_xlsx(jira_server, tmp_path):
    result = discover_fields(client(jira_server), "ABC", sample_size=30)
    out = tmp_path / "fields.xlsx"
    write_field_list_xlsx(out, result)
    ws = load_workbook(out).active
    assert ws.title == "ABC fields"
    assert [c.value for c in ws[1]] == [
        "Field name", "Field id", "Custom", "Type", "Issues with a value (of 30 checked)",
        "On create screen for"]
    rows = {r[1]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert len(rows) == len(result.fields)
    assert rows["key"][4] == 30
    assert rows["customfield_10003"] == ("Severity", "customfield_10003", "Yes", "option", 0, "Bug")
    assert rows["summary"][5] == "Bug, Story"
    assert ws.freeze_panes == "A2"


# ------------------------------------------------------------------- UI ---

def test_dialog_scan_and_save(jira_server, tmp_path, monkeypatch):
    from PySide6.QtCore import QSettings, Qt
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

    from jira_exporter import app as appmod
    from test_ui import wait_until

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)
    warned = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: warned.append(a))
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()

    win.login_action.trigger()
    lp = win.login
    lp.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    lp.username.setText("alice")
    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    assert win.stack.currentWidget() is win.start
    win.jql_action.trigger()

    ep = win.export
    ep.config_path.setText(str(ROOT / "config" / "fields.json"))
    win.discover_action.trigger()
    dlg = ep.fields_dialog
    assert dlg is not None and not dlg.save_btn.isEnabled()

    dlg.project.setText("NOPE")
    dlg.scan_btn.click()
    wait_until(qapp, lambda: "Scan failed" in dlg.status.text())
    assert "was not found" in dlg.status.text() and dlg.scan_btn.isEnabled()

    dlg.project.setText("ABC")
    dlg.sample.setValue(100)
    dlg.scan_btn.click()
    wait_until(qapp, lambda: dlg.result is not None)
    assert "Alpha Beta (ABC)" in dlg.status.text()

    def shown_ids():
        return [dlg.table.item(r, 2).text() for r in range(dlg.table.rowCount())]

    assert shown_ids()[0] == "key"
    assert "customfield_10003" in shown_ids() and "customfield_10099" not in shown_ids()
    dlg.show_all.setChecked(True)
    assert "customfield_10099" in shown_ids()
    dlg.show_all.setChecked(False)

    # Fields with values are ticked; Severity (only on a create screen) is not.
    ticked = {f.id for f in dlg.selected_fields()}
    assert {"issuekey", "summary", "assignee", "customfield_10002"} <= ticked
    assert "customfield_10003" not in ticked
    dlg.filter.setText("story")
    hidden = [dlg.table.isRowHidden(r) for r in range(dlg.table.rowCount())]
    assert hidden.count(False) == 1
    dlg.filter.setText("")

    dlg.clear_btn.click()
    dlg.save_btn.click()
    assert warned  # nothing ticked
    for row, f in enumerate(dlg._rows):
        if f.id in ("issuekey", "summary", "customfield_10003"):
            dlg.table.item(row, 0).setCheckState(Qt.Checked)

    target = tmp_path / "abc_fields"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(target), ""))
    dlg.save_btn.click()
    saved = tmp_path / "abc_fields.json"
    assert json.loads(saved.read_text(encoding="utf-8"))["fields"] == [
        {"id": "key", "header": "Issue Key"}, {"id": "summary", "header": "Summary"},
        {"id": "customfield_10003", "header": "Severity"}]
    assert ep.config_path.text() == str(saved)

    xlsx = tmp_path / "list"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(xlsx), ""))
    dlg.xlsx_btn.click()
    assert (tmp_path / "list.xlsx").is_file()
    dlg.reject()

    # The export screen now uses the saved file.
    out = tmp_path / "abc.xlsx"
    ep.jql.setPlainText("project = ABC")
    ep.out_path.setText(str(out))
    ep.export_btn.click()
    wait_until(qapp, lambda: "Done: 250 issues" in ep.log.toPlainText())
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["Issue Key", "Summary", "Severity"]


def test_dialog_close_cancels_scan(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter.fields_dialog import FieldsDialog

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    QApplication.instance() or QApplication([])
    settings = QSettings("JiraDesktopExporter", "JiraDesktopExporter")
    dlg = FieldsDialog(client(jira_server), settings, "")
    dlg.project.setText("ABC")
    dlg.sample.setValue(5000)
    dlg.scan_btn.click()
    dlg.reject()  # must stop the worker thread without hanging or crashing
    assert dlg._thread is None and dlg.result is None
    assert settings.value("discover_project") == "ABC"


def test_suggested_fields_skip_raw_data_and_placeholders(jira_server, monkeypatch):
    c = client(jira_server)
    real = c.search

    def search(*a, **k):
        for issue in real(*a, **k):
            yield dict(issue, fields=dict(issue["fields"], workratio=-1,
                                          comment={"comments": [{"body": "x"}], "total": 1}))
    monkeypatch.setattr(c, "search", search)
    f = _by_id(discover_fields(c, "ABC", sample_size=10))
    assert "workratio" not in f or f["workratio"].used_in == 0
    assert f["comment"].used_in == 10 and not f["comment"].suggested
    assert f["summary"].suggested and f["issuekey"].suggested


def test_no_sample_and_no_createmeta_says_why(jira_server):
    jira_server.createmeta_status = 403
    result = discover_fields(client(jira_server), "ABC", sample_size=0)
    assert [x.id for x in result.relevant] == ["issuekey"]
    assert any("No issues were checked" in w for w in result.warnings)


def test_write_config_with_unreadable_template(tmp_path):
    template = tmp_path / "ansi.json"
    template.write_bytes('{"fields": [{"name": "R\xe9sum\xe9"}]}'.encode("cp1252"))
    out = tmp_path / "out.json"
    write_config(out, [FieldUsage("summary", "Summary")], template=template)
    assert load_config(out).fields[0].id == "summary"


def _dialog(srv, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter.fields_dialog import FieldsDialog

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    dlg = FieldsDialog(client(srv), QSettings("JiraDesktopExporter", "JiraDesktopExporter"), "")
    dlg.show()
    return qapp, dlg


def _scan(qapp, dlg, key="ABC"):
    from test_ui import wait_until
    dlg.project.setText(key)
    before = dlg.result
    dlg.scan_btn.click()
    wait_until(qapp, lambda: dlg._worker is None)
    return dlg.result is not before


def test_enter_in_filter_does_not_rescan(jira_server, tmp_path):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    qapp, dlg = _dialog(jira_server, tmp_path)
    assert _scan(qapp, dlg)
    searches = len(jira_server.requests)
    dlg._tick(lambda f: f.id == "customfield_10003")
    QTest.keyClick(dlg.filter, Qt.Key_Return)
    QTest.keyClick(dlg.table, Qt.Key_Return)
    for _ in range(20):
        qapp.processEvents()
    assert len(jira_server.requests) == searches and dlg._worker is None
    assert [f.id for f in dlg.selected_fields()] == ["customfield_10003"]
    dlg.reject()


def test_hidden_ticks_are_kept_and_saved(jira_server, tmp_path):
    from PySide6.QtCore import Qt

    qapp, dlg = _dialog(jira_server, tmp_path)
    _scan(qapp, dlg)
    dlg.show_all.setChecked(True)
    row = [f.id for f in dlg._rows].index("customfield_10099")
    dlg.table.item(row, 0).setCheckState(Qt.Checked)
    dlg.show_all.setChecked(False)
    assert "customfield_10099" in {f.id for f in dlg.selected_fields()}
    dlg.show_all.setChecked(True)
    row = [f.id for f in dlg._rows].index("customfield_10099")
    assert dlg.table.item(row, 0).checkState() == Qt.Checked
    dlg.reject()


def test_failed_rescan_keeps_previous_result_usable(jira_server, tmp_path):
    qapp, dlg = _dialog(jira_server, tmp_path)
    _scan(qapp, dlg)
    _scan(qapp, dlg, "ABD")
    assert "was not found" in dlg.status.text() and "previous scan of ABC" in dlg.status.text()
    assert dlg.save_btn.isEnabled() and dlg.result.project_key == "ABC"
    dlg.reject()


def test_status_shows_special_characters_as_typed(jira_server, tmp_path, monkeypatch):
    import conftest
    monkeypatch.setitem(conftest.PROJECT, "name", "R&D <Platform>")
    qapp, dlg = _dialog(jira_server, tmp_path)
    _scan(qapp, dlg)
    assert dlg.status.text().startswith("R&D <Platform> (ABC)")
    dlg.reject()


def test_closing_during_a_slow_scan_does_not_freeze(jira_server, tmp_path):
    import time
    from test_ui import wait_until

    jira_server.search_delay = 1.5
    qapp, dlg = _dialog(jira_server, tmp_path)
    dlg.project.setText("ABC")
    dlg.scan_btn.click()
    wait_until(qapp, lambda: "Checking which fields" in dlg.status.text())
    start = time.monotonic()
    dlg.reject()
    assert time.monotonic() - start < 0.5
    assert dlg.result is None
    from jira_exporter import fields_dialog
    wait_until(qapp, lambda: not fields_dialog._RUNNING, timeout=10)
    for _ in range(20):
        qapp.processEvents()
    assert dlg.result is None  # the late result was ignored


def test_tick_all_ticks_listed_rows(jira_server, tmp_path):
    qapp, dlg = _dialog(jira_server, tmp_path)
    _scan(qapp, dlg)
    dlg.clear_btn.click()
    dlg.filter.setText("story")
    dlg.tick_all_btn.click()
    assert [f.id for f in dlg.selected_fields()] == ["customfield_10002"]
    dlg.filter.setText("")
    dlg.tick_all_btn.click()
    assert {f.id for f in dlg.selected_fields()} == {f.id for f in dlg.result.relevant}
    dlg.show_all.setChecked(True)
    dlg.tick_all_btn.click()
    assert len(dlg.selected_fields()) == len(dlg.result.fields)
    dlg.reject()


# ------------------------------------------------------- one sample issue ---

@pytest.mark.parametrize("cloud", [False, True])
def test_discover_from_one_issue(jira_server, cloud_server, cloud):
    from jira_exporter.field_discovery import discover_from_issue
    srv = cloud_server if cloud else jira_server
    result = discover_from_issue(client(srv, CLOUD if cloud else "datacenter"), "abc-3")
    api = "/rest/api/3" if cloud else "/rest/api/2"
    assert [p for p, _ in srv.requests] == [f"{api}/issue/ABC-3?expand=names%2Cschema"]
    assert (result.project_key, result.project_name, result.source_issue, result.sampled) == \
        ("ABC", "Alpha Beta", "ABC-3", 1)
    by = _by_id(result)
    assert by["issuekey"].used_in == 1 and result.fields[0].is_key
    assert by["summary"].used_in == 1 and by["summary"].name == "Summary"
    assert by["labels"].type == "array"
    assert by["customfield_10002"].name == "Story Points"
    # Empty on this issue: listed, but not used or ticked.
    assert by["assignee"].used_in == 0  # ABC-3 has no assignee
    assert by["customfield_10003"].used_in == 0 and not by["customfield_10003"].suggested
    assert by["customfield_10003"].custom and by["customfield_10003"].name == "Severity"
    assert by["workratio"].used_in == 0  # -1 = no estimate
    assert "one issue, ABC-3" in result.warnings[0]


def test_discover_from_project_uses_latest_issue(jira_server):
    from jira_exporter.field_discovery import discover_from_issue
    result = discover_from_issue(client(jira_server), "ABC")
    assert result.source_issue == "ABC-1"
    paths = [p for p, _ in jira_server.requests]
    assert paths == ["/rest/api/2/search", "/rest/api/2/issue/ABC-1?expand=names%2Cschema"]
    assert jira_server.requests[0][1]["maxResults"] == 1


def test_discover_from_issue_errors(jira_server):
    from jira_exporter.field_discovery import discover_from_issue
    with pytest.raises(JiraError, match="Issue 'ABC-999' was not found"):
        discover_from_issue(client(jira_server), "ABC-999")
    with pytest.raises(JiraError, match="Project 'NOPE' was not found"):
        discover_from_issue(client(jira_server), "NOPE")


def test_dialog_one_issue_mode(jira_server, tmp_path):
    qapp, dlg = _dialog(jira_server, tmp_path)
    assert dlg.sample.isVisibleTo(dlg)
    dlg.mode.setCurrentIndex(1)
    assert not dlg.sample.isVisibleTo(dlg) and "Issue key" in dlg.project.placeholderText()
    assert _scan(qapp, dlg, "ABC-3")
    assert dlg.result.source_issue == "ABC-3"
    assert "fields have values on ABC-3" in dlg.status.text()
    assert len(jira_server.requests) == 1
    ids = {f.id for f in dlg.selected_fields()}
    assert {"issuekey", "summary", "status"} <= ids and "customfield_10003" not in ids
    assert dlg.settings.value("discover_one_issue", type=bool) is True
    dlg.reject()
