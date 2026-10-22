"""Jira answering too slowly: retry reads, and ask for smaller search pages."""
import pytest

from jira_exporter.jira_client import (
    CLOUD, Credentials, JiraClient, JiraError, JiraTimeout)


def client(srv, deployment="datacenter", timeout=0.3):
    return JiraClient(Credentials(f"http://127.0.0.1:{srv.server_address[1]}", "alice", "secret",
                                  deployment=deployment), timeout=timeout)


@pytest.fixture(autouse=True)
def quick_retries(monkeypatch):
    monkeypatch.setattr(JiraClient, "TIMEOUT_RETRY_WAITS", (0.0, 0.0))


def test_default_timeout_is_generous():
    c = JiraClient(Credentials("https://jira.example.com", "a", "b"))
    assert c.timeout == 120


@pytest.mark.parametrize("cloud", [False, True])
def test_search_shrinks_pages_that_time_out(jira_server, cloud_server, cloud):
    srv = cloud_server if cloud else jira_server
    srv.slow_above = 30  # pages of more than 30 issues take too long
    c = client(srv, CLOUD if cloud else "datacenter")
    notes = []
    c.notify_waits(None, notes.append)
    issues = list(c.search("project = ABC", ["summary"], page_size=100))
    assert len(issues) == 250
    assert [i["key"] for i in issues] == [f"ABC-{n}" for n in range(1, 251)]  # no gaps/dups
    assert srv.slow_pages == [100, 50]  # then 25 at a time
    assert any("asking for 50 at a time" in n for n in notes)
    assert any("asking for 25 at a time" in n for n in notes)


def test_search_gives_up_below_the_smallest_page(jira_server):
    jira_server.slow_above = 0  # every page is too slow
    with pytest.raises(JiraTimeout) as err:
        list(client(jira_server).search("project = ABC", ["summary"], page_size=40))
    # 40 -> 20 -> 10, then 10 is retried twice before giving up.
    assert jira_server.slow_pages == [40, 20, 10, 10, 10]
    msg = str(err.value)
    assert "tried 3 times" in msg and "max_results_per_page" in msg


def test_get_is_retried_after_a_timeout(jira_server):
    jira_server.slow_gets = 2
    c = client(jira_server)
    notes = []
    c.notify_waits(None, notes.append)
    assert c.myself()["name"] == "alice"
    assert len(notes) == 2 and "trying again" in notes[0]


def test_get_gives_up_after_retries(jira_server):
    jira_server.slow_gets = 3
    with pytest.raises(JiraError, match=r"didn.t answer within .*\(tried 3 times\)"):
        client(jira_server).fields()


def test_post_that_isnt_a_search_is_not_retried(jira_server, monkeypatch):
    # Creating a sprint must not be repeated: the first try may have worked.
    c = client(jira_server)
    calls = []
    import requests

    def slow(*a, **k):
        calls.append(a)
        raise requests.exceptions.ReadTimeout("read timed out")

    monkeypatch.setattr(c.session, "request", slow)
    with pytest.raises(JiraTimeout):
        c.create_sprint("S 01", 1)
    assert len(calls) == 1


def test_export_keeps_going_after_slow_pages(jira_server, tmp_path):
    from pathlib import Path
    from openpyxl import load_workbook
    from jira_exporter.config import load_config, resolve_fields
    from jira_exporter.exporter import write_xlsx
    root = Path(__file__).resolve().parents[1]
    cfg = load_config(root / "config" / "fields.json")
    jira_server.slow_above = 60
    c = client(jira_server)
    fields, _ = resolve_fields(cfg.fields, c.fields())
    out = tmp_path / "x.xlsx"
    count = write_xlsx(out, fields, c.search("project = ABC", [f.id for f in fields
                                                                 if f.id != "key"], 100))
    assert count == 250 and load_workbook(out).active.max_row == 251
