"""Burndown Chart: replaying Jira's changes, the Excel chart and the dialog."""
import os
from datetime import date, datetime

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook

from conftest import BURNDOWNS, _chg, local_ms
from jira_exporter.burndown import COLUMNS, BurndownError, build_burndown, write_burndown_xlsx
from jira_exporter.jira_client import JiraError
from test_sprint import ROWS, client, write_rows


def test_build_burndown_replays_changes_day_by_day():
    c = build_burndown(BURNDOWNS[(1, 101)])
    assert (c.start, c.end, c.stop) == (datetime(2026, 1, 5, 9), datetime(2026, 1, 9, 9),
                                        datetime(2026, 1, 9, 8))
    assert c.completed and c.start_remaining == 16 and c.remaining == 5
    assert [d.values()[:2] + d.values()[3:] for d in c.days] == [
        [date(2026, 1, 5), 16, 0, 0],
        [date(2026, 1, 6), 11, 5, 0],   # ABC-1 done
        [date(2026, 1, 7), 10, 0, -1],  # ABC-4 added (+2), ABC-3 re-estimated 8 -> 5
        [date(2026, 1, 8), 5, 3, -2],   # ABC-2 done, reopened, done; ABC-4 removed
        [date(2026, 1, 9), 5, 0, 0],
    ]
    # The guideline falls from 16 to 0 over the four planned days (96 hours).
    assert [round(d.guideline, 2) for d in c.days] == [13.5, 9.5, 5.5, 1.5, 0.17]
    assert c.steps == [
        (datetime(2026, 1, 5, 9), 16), (datetime(2026, 1, 6, 10), 11),
        (datetime(2026, 1, 7, 11), 13), (datetime(2026, 1, 7, 15), 10),
        (datetime(2026, 1, 8, 9), 7), (datetime(2026, 1, 8, 12), 10),
        (datetime(2026, 1, 8, 13), 7), (datetime(2026, 1, 8, 16), 5)]


def test_active_sprint_runs_to_now():
    data = {"startTime": local_ms(2026, 3, 2, 9), "endTime": local_ms(2026, 3, 16, 9),
            "now": local_ms(2026, 3, 4, 12),
            "changes": {str(local_ms(2026, 3, 1)): [_chg("A-1", True, 3)],
                        str(local_ms(2026, 3, 20)): [_chg("A-1", done=True)]}}
    c = build_burndown(data)
    assert not c.completed and c.stop == datetime(2026, 3, 4, 12)
    assert [d.day.day for d in c.days] == [2, 3, 4] and c.remaining == 3


def test_issues_without_estimates_and_odd_entries():
    data = {"startTime": local_ms(2026, 3, 2), "endTime": local_ms(2026, 3, 3),
            "completeTime": local_ms(2026, 3, 3),
            "changes": {"bad": [_chg("A-9", True, 9)],
                        str(local_ms(2026, 3, 1)): [_chg("A-1", True), {"no": "key"}, "x",
                                                    {"key": "A-2", "added": True,
                                                     "statC": {"newValue": "n/a"}}]}}
    c = build_burndown(data)
    assert c.start_remaining == 0 and c.remaining == 0


def test_sprint_not_started():
    with pytest.raises(BurndownError, match="may not have started"):
        build_burndown({"changes": {}})


def test_client_scope_change_burndown(jira_server):
    assert client(jira_server).scope_change_burndown(1, 101) == BURNDOWNS[(1, 101)]
    with pytest.raises(JiraError, match="no burndown data for sprint 103 on board 1"):
        client(jira_server).scope_change_burndown(1, 103)


def test_write_burndown_xlsx(tmp_path):
    path = tmp_path / "b.xlsx"
    write_burndown_xlsx(path, "Sprint 01", build_burndown(BURNDOWNS[(1, 101)]))
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.cell(2, 1).value == datetime(2026, 1, 5)
    assert [c.value for c in ws[3]][1:] == [11, 9.5, 5, 0]
    assert ws.cell(6, 3).value == 0.2
    assert ws.max_row == 6 and len(ws._charts) == 1


def test_burndown_sprints_active_first_then_newest():
    from jira_exporter.burndown_dialog import burndown_sprints
    sprints = [{"id": 1, "state": "closed", "completeDate": "2026-01-01"},
               {"id": 2, "state": "closed", "completeDate": "2026-02-01", "originBoardId": 1},
               {"id": 3, "state": "future"}, {"id": 4, "state": "active"},
               {"id": 5, "state": "closed", "originBoardId": 9}]
    assert [s["id"] for s in burndown_sprints(sprints, 1)] == [4, 2, 1]


# ------------------------------------------------------------------ dialog ---

@pytest.fixture
def dialog(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter.burndown_dialog import BurndownDialog
    from test_ui import wait_until

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    settings = QSettings("JiraDesktopExporter", "JiraDesktopExporter")
    settings.setValue("sprint_config", str(write_rows(tmp_path / "teams.xlsx", ROWS)))
    dlg = BurndownDialog(client(jira_server), settings, None)
    dlg.show()
    wait_until(qapp, lambda: dlg._worker is None)
    yield qapp, dlg
    dlg.reject()


def _wait(qapp, dlg):
    from test_ui import wait_until
    wait_until(qapp, lambda: dlg._worker is None)


def test_dialog_lists_sprints_and_charts_one(dialog):
    qapp, dlg = dialog
    assert [dlg.sprint.itemText(i) for i in range(dlg.sprint.count())] == [
        "Team Alpha Sprint 03 (active)", "Team Alpha Sprint 02", "Team Alpha Sprint 01"]
    assert not dlg.export_btn.isEnabled()
    dlg.sprint.setCurrentIndex(2)
    dlg.run_btn.click()
    _wait(qapp, dlg)
    assert dlg.status.text() == ("Team Alpha Sprint 01: 16 story points at the start, "
                                 "5 remaining when it closed.")
    assert dlg.table.rowCount() == 5
    assert [dlg.table.item(1, c).text() for c in range(5)] == \
        ["2026-01-06", "11", "9.5", "5", "0"]
    guideline, remaining = dlg.chart.series()
    assert remaining.name() == "Remaining" and guideline.count() == 2
    # Steps: each change after the first adds two points, plus one at the close.
    assert remaining.count() == 1 + 2 * 7 + 1
    assert remaining.at(remaining.count() - 1).y() == 5
    assert dlg.chart.title() == "Team Alpha Sprint 01 burndown"
    assert dlg.export_btn.isEnabled() and dlg.image_btn.isEnabled()


def test_dialog_sprint_without_data(dialog):
    qapp, dlg = dialog
    dlg.sprint.setCurrentIndex(0)
    dlg.run_btn.click()
    _wait(qapp, dlg)
    assert "no burndown data for sprint 103" in dlg.status.text()
    assert not dlg.export_btn.isEnabled()


def test_dialog_team_without_sprints(dialog):
    qapp, dlg = dialog
    dlg.team.setCurrentText("Team Beta")
    _wait(qapp, dlg)
    assert dlg.sprint.count() == 0 and not dlg.run_btn.isEnabled()
    assert "No active or closed sprints" in dlg.status.text()


def test_dialog_exports_excel_and_image(dialog, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    qapp, dlg = dialog
    dlg.sprint.setCurrentIndex(2)
    dlg.run_btn.click()
    _wait(qapp, dlg)
    answers = iter([str(tmp_path / "b"), str(tmp_path / "chart")])
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (next(answers), ""))
    dlg.export_btn.click()
    dlg.image_btn.click()
    assert load_workbook(tmp_path / "b.xlsx").active.max_row == 6
    assert (tmp_path / "chart.png").stat().st_size > 0


def test_menu_opens_burndown_only_when_logged_in(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter import app as appmod
    from test_ui import wait_until

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    assert not win.burndown_action.isEnabled()
    win.burndown_action.trigger()
    assert win.burndown_dialog is None

    win.login_action.trigger()
    win.login.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    win.login.username.setText("alice")
    win.login.password.setText("secret")
    win.login.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    win.burndown_action.trigger()
    dlg = win.burndown_dialog
    assert dlg is not None and dlg.isVisible()
    win.burndown_action.trigger()
    assert win.burndown_dialog is dlg

    win.logout_action.trigger()
    wait_until(qapp, lambda: win.burndown_dialog is None)
    win.close()
