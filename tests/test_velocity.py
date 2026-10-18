"""Velocity Chart: the average, the Excel chart and the dialog."""
import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook

from conftest import SPRINT_REPORTS
from jira_exporter.sprint_report import summarize
from jira_exporter.velocity import COLUMNS, average_velocity, write_velocity_xlsx
from test_sprint import ROWS, client, write_rows

ROWS_ALPHA = [
    summarize("Team Alpha", {"id": 101, "name": "Team Alpha Sprint 01"},
              SPRINT_REPORTS[(1, 101)]),
    summarize("Team Alpha", {"id": 102, "name": "Team Alpha Sprint 02"},
              SPRINT_REPORTS[(1, 102)]),
]


def test_average_velocity_is_mean_completed_points():
    assert average_velocity(ROWS_ALPHA) == pytest.approx((13 + 8) / 2)
    assert average_velocity([]) is None


def test_write_velocity_xlsx_has_table_average_and_chart(tmp_path):
    path = tmp_path / "v.xlsx"
    write_velocity_xlsx(path, "Team Alpha", ROWS_ALPHA)
    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == COLUMNS
    assert [c.value for c in ws[2]][:3] == ["Team Alpha Sprint 01", 21, 13]
    assert [c.value for c in ws[3]][:3] == ["Team Alpha Sprint 02", 8, 8]
    assert ws.cell(2, 4).number_format == "0%"
    assert (ws.cell(5, 1).value, ws.cell(5, 3).value) == ("Average velocity", 10.5)
    assert len(ws._charts) == 1


def test_write_velocity_xlsx_without_sprints(tmp_path):
    path = tmp_path / "v.xlsx"
    write_velocity_xlsx(path, "Team Alpha", [])
    ws = load_workbook(path).active
    assert ws.cell(3, 1).value == "Average velocity" and ws.cell(3, 3).value is None


@pytest.fixture
def dialog(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter.velocity_dialog import VelocityDialog

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    settings = QSettings("JiraDesktopExporter", "JiraDesktopExporter")
    settings.setValue("sprint_config", str(write_rows(tmp_path / "teams.xlsx", ROWS)))
    dlg = VelocityDialog(client(jira_server), settings, None)
    dlg.show()
    yield qapp, dlg
    dlg.reject()


def _run(qapp, dlg):
    from test_ui import wait_until
    dlg.run_btn.click()
    wait_until(qapp, lambda: dlg._worker is None)


def test_short_names():
    from jira_exporter.velocity_dialog import short_names
    assert short_names("Team A", ["Team A Sprint 1", "Team A 2"]) == ["Sprint 1", "2"]
    assert short_names("Team A", ["Team A Sprint 1", "Other 2"]) == ["Team A Sprint 1", "Other 2"]
    assert short_names("Team A", ["Team A "]) == ["Team A "]


def test_dialog_charts_one_team(dialog):
    from PySide6.QtCore import Qt
    qapp, dlg = dialog
    assert [dlg.team.itemText(i) for i in range(dlg.team.count())] == \
        ["Team Alpha", "Team Beta", "Team Kanban", "Team Lost"]
    assert dlg.count.value() == 7
    assert not dlg.export_btn.isEnabled() and not dlg.image_btn.isEnabled()
    _run(qapp, dlg)
    table = [[dlg.table.item(r, c).text() for c in range(4)]
             for r in range(dlg.table.rowCount())]
    assert table == [["Team Alpha Sprint 01", "21", "13", "62%"],
                     ["Team Alpha Sprint 02", "8", "8", "100%"]]
    (series,) = dlg.chart.series()
    committed, completed = series.barSets()
    assert [committed.at(i) for i in range(2)] == [21, 8]
    assert [completed.at(i) for i in range(2)] == [13, 8]
    assert dlg.chart.title() == "Team Alpha velocity"
    assert dlg.chart.axes(Qt.Horizontal)[0].categories() == ["Sprint 01", "Sprint 02"]
    assert "Average velocity: 10.5 story points per sprint" in dlg.status.text()
    assert dlg.export_btn.isEnabled() and dlg.image_btn.isEnabled()
    assert dlg.settings.value("velocity_team") == "Team Alpha"


def test_dialog_rerun_replaces_the_chart(dialog):
    qapp, dlg = dialog
    _run(qapp, dlg)
    dlg.count.setValue(1)
    _run(qapp, dlg)
    (series,) = dlg.chart.series()
    assert series.barSets()[1].count() == 1
    assert len(dlg.chart.axes()) == 2


def test_dialog_team_without_sprints(dialog):
    qapp, dlg = dialog
    dlg.team.setCurrentText("Team Beta")
    _run(qapp, dlg)
    assert "No closed sprints to chart." in dlg.status.text()
    assert dlg.table.rowCount() == 0 and not dlg.export_btn.isEnabled()


def test_dialog_exports_excel_and_image(dialog, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog
    qapp, dlg = dialog
    _run(qapp, dlg)
    answers = iter([str(tmp_path / "v"), str(tmp_path / "chart")])
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (next(answers), ""))
    dlg.export_btn.click()
    dlg.image_btn.click()
    assert load_workbook(tmp_path / "v.xlsx").active.cell(2, 1).value == "Team Alpha Sprint 01"
    assert (tmp_path / "chart.png").stat().st_size > 0


def test_menu_opens_velocity_only_when_logged_in(jira_server, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QApplication

    from jira_exporter import app as appmod
    from test_ui import wait_until

    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    assert not win.velocity_action.isEnabled()
    win.velocity_action.trigger()
    assert win.velocity_dialog is None

    win.login_action.trigger()
    win.login.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    win.login.username.setText("alice")
    win.login.password.setText("secret")
    win.login.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    win.velocity_action.trigger()
    dlg = win.velocity_dialog
    assert dlg is not None and dlg.isVisible()
    win.velocity_action.trigger()
    assert win.velocity_dialog is dlg

    win.logout_action.trigger()
    wait_until(qapp, lambda: win.velocity_dialog is None)
    win.close()
