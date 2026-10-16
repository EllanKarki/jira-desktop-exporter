from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import FIELDS
from jira_exporter.config import ConfigError, FieldSpec, load_config, resolve_fields
from conftest import ADF_DESC
from jira_exporter.exporter import adf_to_text, format_value, write_xlsx
from jira_exporter.jira_client import (
    CLOUD, DATA_CENTER, Credentials, JiraClient, JiraError, guess_deployment,
)

ROOT = Path(__file__).resolve().parents[1]


def url(srv):
    return f"http://127.0.0.1:{srv.server_address[1]}"


def test_default_config_loads():
    cfg = load_config(ROOT / "config" / "fields.json")
    assert cfg.fields[0].id == "key"
    assert cfg.page_size == 100


def test_bad_config(tmp_path):
    p = tmp_path / "f.json"
    p.write_text('{"fields": []}')
    with pytest.raises(ConfigError):
        load_config(p)
    p.write_text("{not json")
    with pytest.raises(ConfigError):
        load_config(p)


def test_resolve_by_name_and_id():
    specs = [FieldSpec("key", None), FieldSpec(None, "story points", "SP"),
             FieldSpec("status", None), FieldSpec(None, "Nope")]
    fields, problems = resolve_fields(specs, FIELDS)
    assert [f.id for f in fields] == ["key", "customfield_10002", "status"]
    assert fields[1].header == "SP"
    assert problems == ["Field name 'Nope' was not found in Jira."]


def test_format_value():
    assert format_value(None) == ""
    assert format_value({"displayName": "X", "name": "x"}) == "X"
    assert format_value([{"name": "1.0"}, {"name": "2.0"}]) == "1.0, 2.0"
    assert format_value({"value": "Parent", "child": {"value": "Kid"}}) == "Parent - Kid"
    assert format_value("2026-01-02T10:15:30.000+0530", "datetime") == datetime(2026, 1, 2, 10, 15, 30)
    assert format_value("2026-02-01", "date") == date(2026, 2, 1)
    assert format_value(3.5) == 3.5


def test_login_password_and_token(jira_server):
    assert JiraClient(Credentials(url(jira_server), "alice", "secret")).myself()["name"] == "alice"
    assert JiraClient(Credentials(url(jira_server), "", "good-token", use_token=True)).myself()
    with pytest.raises(JiraError, match="Login failed"):
        JiraClient(Credentials(url(jira_server), "alice", "wrong")).myself()


def test_unreachable():
    with pytest.raises(JiraError, match="Could not reach"):
        JiraClient(Credentials("http://127.0.0.1:1", "a", "b"), timeout=2).myself()


def test_bad_jql(jira_server):
    c = JiraClient(Credentials(url(jira_server), "alice", "secret"))
    with pytest.raises(JiraError, match="Error in the JQL"):
        list(c.search("bad", ["summary"]))


def test_full_export_paginates(jira_server, tmp_path):
    c = JiraClient(Credentials(url(jira_server), "alice", "secret"))
    cfg = load_config(ROOT / "config" / "fields.json")
    fields, _ = resolve_fields(
        [FieldSpec("key", None, "Issue Key"), FieldSpec(None, "Summary"),
         FieldSpec(None, "Assignee"), FieldSpec(None, "Created"), FieldSpec(None, "Labels")],
        c.fields(),
    )
    out = tmp_path / "out.xlsx"
    n = write_xlsx(out, fields, c.search("project = ABC", [f.id for f in fields[1:]], 100))
    assert n == 250
    assert [b["startAt"] for _, b in jira_server.requests] == [0, 100, 200]
    ws = load_workbook(out).active
    assert [c.value for c in ws[1]] == ["Issue Key", "Summary", "Assignee", "Created", "Labels"]
    assert ws["A2"].value == "ABC-1"
    assert ws["C2"].value == "Dev One"
    assert ws["C4"].value is None  # unassigned
    assert ws["D2"].value == datetime(2026, 1, 2, 10, 15, 30)
    assert ws["E2"].value == "a, b"
    assert ws.max_row == 251


def test_guess_deployment():
    assert guess_deployment("https://acme.atlassian.net") == CLOUD
    assert guess_deployment("acme.atlassian.net/") == CLOUD
    assert guess_deployment("https://jira.acme.com") == DATA_CENTER
    assert guess_deployment("https://atlassian.net.evil.com") == DATA_CENTER


def test_adf_to_text():
    assert adf_to_text(ADF_DESC).strip() == "Hello world\n- one\n- two"
    assert format_value(ADF_DESC) == "Hello world\n- one\n- two"


def test_cloud_login_uses_v3_and_basic_auth(cloud_server):
    # Cloud uses email + API token as Basic auth; the fake accepts "alice:secret".
    c = JiraClient(Credentials(url(cloud_server), "alice", "secret", use_token=True,
                               deployment=CLOUD))
    assert c.api == "/rest/api/3"
    assert c.myself()["displayName"] == "Alice Cloud"
    with pytest.raises(JiraError, match="email and API token"):
        JiraClient(Credentials(url(cloud_server), "alice", "bad", deployment=CLOUD)).myself()


def test_cloud_export_uses_token_pagination(cloud_server, tmp_path):
    c = JiraClient(Credentials(url(cloud_server), "alice", "secret", deployment=CLOUD))
    fields, _ = resolve_fields(
        [FieldSpec("key", None), FieldSpec(None, "Summary"), FieldSpec(None, "Description")],
        c.fields(),
    )
    seen = []
    out = tmp_path / "cloud.xlsx"
    n = write_xlsx(out, fields, c.search("project = ABC", ["summary", "description"], 100,
                                         progress=lambda d, t: seen.append((d, t))))
    assert n == 250
    paths = [p for p, _ in cloud_server.requests]
    assert paths == ["/rest/api/3/search/approximate-count"] + ["/rest/api/3/search/jql"] * 3
    tokens = [b.get("nextPageToken") for p, b in cloud_server.requests[1:]]
    assert tokens == [None, "100", "200"]
    assert seen == [(100, 250), (200, 250), (250, 250)]
    ws = load_workbook(out).active
    assert ws["A2"].value == "ABC-1"
    assert ws["C2"].value == "Hello world\n- one\n- two"


def test_cloud_bad_jql(cloud_server):
    c = JiraClient(Credentials(url(cloud_server), "alice", "secret", deployment=CLOUD))
    with pytest.raises(JiraError, match="Error in the JQL"):
        list(c.search("bad", ["summary"]))
