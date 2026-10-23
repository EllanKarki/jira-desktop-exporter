"""Export choice: replace the Excel file, or add to it and update issues already there."""
from datetime import datetime

from openpyxl import Workbook, load_workbook

from conftest import ISSUES
from jira_exporter.config import ResolvedField
from jira_exporter.exporter import append_xlsx, write_xlsx

FIELDS = [ResolvedField("key", "Issue Key"), ResolvedField("summary", "Summary"),
          ResolvedField("created", "Created", "datetime")]


def _issue(key, summary, created="2026-01-02T10:15:30.000+0530"):
    return {"key": key, "fields": {"summary": summary, "created": created}}


def _rows(path):
    ws = load_workbook(path)["Issues"]
    return [[c.value for c in row] for row in ws.iter_rows()]


def test_append_adds_new_and_updates_existing(tmp_path):
    out = tmp_path / "e.xlsx"
    write_xlsx(out, FIELDS, [_issue("ABC-1", "one"), _issue("ABC-2", "two")])
    result = append_xlsx(out, FIELDS, [_issue("ABC-2", "two, edited"), _issue("ABC-3", "three")])
    assert (result.added, result.updated, result.count) == (1, 1, 2)
    rows = _rows(out)
    assert rows[0] == ["Issue Key", "Summary", "Created"]
    assert [r[:2] for r in rows[1:]] == [["ABC-1", "one"], ["ABC-2", "two, edited"],
                                         ["ABC-3", "three"]]
    assert rows[3][2] == datetime(2026, 1, 2, 10, 15, 30)
    ws = load_workbook(out)["Issues"]
    assert ws["C4"].number_format == "yyyy-mm-dd hh:mm"
    assert ws.auto_filter.ref == "A1:C4" and ws.freeze_panes == "A2"


def test_append_matches_columns_by_header(tmp_path):
    out = tmp_path / "e.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append(["Summary", "Issue Key", "My notes"])
    ws.append(["old", "ABC-1", "keep me"])
    wb.save(out)
    fields = FIELDS[:2] + [ResolvedField("status", "Status")]
    append_xlsx(out, fields, [{"key": "ABC-1", "fields": {"summary": "new",
                                                          "status": {"name": "Done"}}},
                              {"key": "ABC-9", "fields": {"summary": "nine"}}])
    rows = _rows(out)
    assert rows[0] == ["Summary", "Issue Key", "My notes", "Status"]
    assert rows[1] == ["new", "ABC-1", "keep me", "Done"]  # column the export lacks is kept
    assert rows[2] == ["nine", "ABC-9", None, None]


def test_append_to_missing_file_creates_it(tmp_path):
    out = tmp_path / "new.xlsx"
    result = append_xlsx(out, FIELDS, [_issue("ABC-1", "one")])
    assert (result.added, result.updated) == (1, 0)
    assert _rows(out)[1][:2] == ["ABC-1", "one"]


def test_append_without_key_column_warns(tmp_path):
    out = tmp_path / "e.xlsx"
    fields = FIELDS[1:2]
    write_xlsx(out, fields, [_issue("ABC-1", "one")])
    result = append_xlsx(out, fields, [_issue("ABC-1", "one")])
    assert result.added == 1 and "no issue key column" in result.warning
    assert len(_rows(out)) == 3


def test_failed_fetch_leaves_file_untouched(tmp_path):
    out = tmp_path / "e.xlsx"
    write_xlsx(out, FIELDS, [_issue("ABC-1", "one")])
    before = out.read_bytes()

    def broken():
        yield _issue("ABC-2", "two")
        raise RuntimeError("Jira went away")

    try:
        append_xlsx(out, FIELDS, broken())
    except RuntimeError:
        pass
    assert out.read_bytes() == before


def test_ui_add_mode_updates_instead_of_duplicating(jira_server, tmp_path, monkeypatch):
    from test_ui import QApplication, QMessageBox, QSettings, ROOT, appmod, wait_until
    for fmt in (QSettings.NativeFormat, QSettings.IniFormat):
        QSettings.setPath(fmt, QSettings.UserScope, str(tmp_path))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    monkeypatch.setattr(QMessageBox, "critical", lambda *a, **k: None)
    qapp = QApplication.instance() or QApplication([])
    win = appmod.MainWindow()
    win.login_action.trigger()
    lp = win.login
    lp.url.setText(f"http://127.0.0.1:{jira_server.server_address[1]}")
    lp.username.setText("alice")
    lp.password.setText("secret")
    lp.connect_btn.click()
    wait_until(qapp, lambda: win.export.client is not None)
    win.jql_action.trigger()

    ep = win.export
    ep.jql.setPlainText("project = ABC")
    ep.config_path.setText(str(ROOT / "config" / "fields.json"))
    out = tmp_path / "ui.xlsx"
    ep.out_path.setText(str(out))
    ep.write_mode.setCurrentIndex(ep.write_mode.findData("append"))
    for run in (1, 2):
        ep.log.clear()
        ep.export_btn.click()
        wait_until(qapp, lambda: "Done: 250 issues" in ep.log.toPlainText())
        wait_until(qapp, lambda: ep.export_btn.isEnabled())
    assert "Added 0 new issues and updated 250" in ep.log.toPlainText()
    assert load_workbook(out).active.max_row == len(ISSUES) + 1
    assert QSettings("JiraDesktopExporter", "JiraDesktopExporter").value("write_mode") == "append"
    win.close()
