"""Turning Jira issues into rows and writing them to an .xlsx file."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jira_exporter.config import ResolvedField

# Excel rejects cells longer than this.
MAX_CELL_LEN = 32767


def format_value(value: Any, schema_type: str | None = None) -> Any:
    """Flatten a Jira field value into something an Excel cell can hold."""
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(format_value(v)) for v in value if v is not None)
    if isinstance(value, dict):
        if value.get("type") == "doc":  # Cloud v3 rich text (Atlassian Document Format)
            # Strip only blank lines so a leading code block keeps its indentation.
            return format_value(adf_to_text(value).strip("\n"))
        for k in ("displayName", "name", "value", "key", "emailAddress"):
            if value.get(k) not in (None, ""):
                text = str(value[k])
                # Cascading select fields carry the child option in "child".
                child = value.get("child")
                if isinstance(child, dict) and child.get("value"):
                    text += f" - {child['value']}"
                return text
        return str(value)
    if isinstance(value, str) and schema_type in ("datetime", "date"):
        parsed = _parse_date(value, schema_type)
        if parsed is not None:
            return parsed
    if isinstance(value, str) and len(value) > MAX_CELL_LEN:
        return value[: MAX_CELL_LEN - 3] + "..."
    return value


_ADF_BLOCKS = {"paragraph", "heading", "blockquote", "codeBlock", "panel", "rule",
               "decisionItem", "mediaSingle"}


def adf_to_text(node: Any) -> str:
    """Flatten an Atlassian Document Format tree into plain text.

    Lists keep their numbering, and nested content is indented under its item.
    Table rows become "a | b | c" lines, and link targets are kept next to their text.
    """
    if isinstance(node, list):
        return _adf_inline_run(node)
    if not isinstance(node, dict):
        return ""
    kind = node.get("type")
    attrs = node.get("attrs") or {}
    content = node.get("content") or []
    if kind == "text":
        return _with_href(node.get("text", ""), _link_href(node))
    if kind == "hardBreak":
        return "\n"
    if kind == "date":
        try:
            ts = int(attrs.get("timestamp")) / 1000
            return datetime.fromtimestamp(ts, timezone.utc).date().isoformat()
        except (TypeError, ValueError, OverflowError, OSError):
            return str(attrs.get("timestamp") or "")
    if kind in ("mention", "emoji", "status", "inlineCard"):
        return str(attrs.get("text") or attrs.get("shortName") or attrs.get("url") or "")
    if kind in ("blockCard", "embedCard"):
        return str(attrs.get("url") or "") + "\n"
    if kind in ("bulletList", "orderedList"):
        first = attrs.get("order", 1) if kind == "orderedList" else None
        items = [c for c in content if isinstance(c, dict)]
        return "".join(
            _list_entry(f"{first + i}. " if first is not None else "- ", adf_to_text(item.get("content") or []))
            for i, item in enumerate(items))
    if kind == "taskList":
        out = []
        for child in content:
            if not isinstance(child, dict):
                continue
            if child.get("type") == "taskList":  # an indented (nested) action item list
                out.append(_indent(adf_to_text(child), "  "))
            else:
                box = "[x] " if (child.get("attrs") or {}).get("state") == "DONE" else "[ ] "
                out.append(_list_entry(box, adf_to_text(child.get("content") or [])))
        return "".join(out)
    if kind == "tableRow":
        cells = [" ".join(adf_to_text(c.get("content") or []).split())
                 for c in content if isinstance(c, dict)]
        return " | ".join(cells) + "\n"
    text = adf_to_text(content)
    if kind in _ADF_BLOCKS:
        text = text.rstrip("\n") + "\n"
    return text


def _link_href(node: Any) -> str | None:
    if not isinstance(node, dict) or node.get("type") != "text":
        return None
    for mark in node.get("marks") or []:
        if isinstance(mark, dict) and mark.get("type") == "link":
            return (mark.get("attrs") or {}).get("href")
    return None


def _with_href(text: str, href: str | None) -> str:
    return f"{text} ({href})" if href and href != text else text


def _adf_inline_run(nodes: list) -> str:
    """Join sibling nodes; text pieces of one link (split by bold etc.) share one URL."""
    out, i = [], 0
    while i < len(nodes):
        href = _link_href(nodes[i])
        if href:
            j = i
            while j < len(nodes) and _link_href(nodes[j]) == href:
                j += 1
            out.append(_with_href("".join(n.get("text", "") for n in nodes[i:j]), href))
            i = j
        else:
            out.append(adf_to_text(nodes[i]))
            i += 1
    return "".join(out)


def _indent(text: str, pad: str) -> str:
    return "".join(pad + line if line.strip() else line for line in text.splitlines(True))


def _list_entry(marker: str, body: str) -> str:
    """One list item: the marker on the first line, later lines indented under it."""
    first, _, rest = body.strip("\n").partition("\n")
    entry = (marker + first).rstrip(" ") + "\n" if not first.strip() else marker + first + "\n"
    if rest:
        entry += _indent(rest + "\n", " " * len(marker))
    return entry


def _parse_date(text: str, schema_type: str):
    try:
        if schema_type == "date":
            return datetime.strptime(text, "%Y-%m-%d").date()
        # Jira sends e.g. 2024-05-01T10:15:30.000+0530; Excel cannot hold tz info.
        return datetime.strptime(text, "%Y-%m-%dT%H:%M:%S.%f%z").replace(tzinfo=None)
    except ValueError:
        return None


def issue_to_row(issue: dict, fields: list[ResolvedField]) -> list[Any]:
    data = issue.get("fields") or {}
    row = []
    for f in fields:
        if f.id == "key":
            row.append(issue.get("key", ""))
        else:
            row.append(format_value(data.get(f.id), f.schema_type))
    return row


def write_xlsx(path: Path, fields: list[ResolvedField], issues: Iterable[dict]) -> int:
    """Write issues to `path`, replacing any file there. Returns the number written."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Issues"
    ws.append([f.header for f in fields])
    count = 0
    for issue in issues:
        ws.append(issue_to_row(issue, fields))
        count += 1
    _finish_sheet(ws, [f.header for f in fields], {f.header: f.schema_type for f in fields})
    wb.save(path)
    return count


@dataclass
class AppendResult:
    added: int = 0
    updated: int = 0
    warning: str = ""

    @property
    def count(self) -> int:
        return self.added + self.updated


def append_xlsx(path: Path, fields: list[ResolvedField], issues: Iterable[dict]) -> AppendResult:
    """Add issues to the existing file at `path` (or create it).

    Columns are matched by header. An issue whose key is already in the file has its
    row overwritten instead of being added twice; new columns go on the right.
    The file is only saved at the end, so a failed fetch leaves it untouched.
    """
    if not path.exists():
        return AppendResult(added=write_xlsx(path, fields, issues))
    wb = load_workbook(path)
    ws = wb["Issues"] if "Issues" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    while headers and headers[-1] in (None, ""):
        headers.pop()
    for f in fields:
        if f.header not in headers:
            headers.append(f.header)
            ws.cell(row=1, column=len(headers), value=f.header)
    col = {h: i + 1 for i, h in enumerate(headers)}

    key_field = next((f for f in fields if f.id == "key"), None)
    result = AppendResult()
    rows: dict[str, int] = {}
    if key_field is None:
        result.warning = ("The fields config has no issue key column, so issues already "
                          "in the file may now appear twice.")
    else:
        kc = col[key_field.header]
        for r in range(2, ws.max_row + 1):
            v = ws.cell(row=r, column=kc).value
            if v not in (None, ""):
                rows.setdefault(str(v), r)

    next_row = ws.max_row + 1
    for issue in issues:
        values = issue_to_row(issue, fields)
        key = issue.get("key", "")
        r = rows.get(key) if key_field is not None else None
        if r is None:
            r, next_row = next_row, next_row + 1
            if key_field is not None and key:
                rows[key] = r
            result.added += 1
        else:
            result.updated += 1
        for f, v in zip(fields, values):
            ws.cell(row=r, column=col[f.header], value=v)

    _finish_sheet(ws, headers, {f.header: f.schema_type for f in fields})
    wb.save(path)
    return result


def _finish_sheet(ws, headers: list, types: dict) -> None:
    """Header style, frozen header row, column widths, date formats and the filter."""
    for cell in ws[1][: len(headers)]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"
    widths = [len(str(h)) for h in headers]
    for row in ws.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        for i, v in enumerate(row):
            if v is not None:
                widths[i] = max(widths[i], min(len(str(v)), 60))
    for i, h in enumerate(headers):
        col = get_column_letter(i + 1)
        ws.column_dimensions[col].width = widths[i] + 2
        fmt = {"datetime": "yyyy-mm-dd hh:mm", "date": "yyyy-mm-dd"}.get(types.get(h))
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=i + 1, max_col=i + 1):
                cell.number_format = fmt
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
