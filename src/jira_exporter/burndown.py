"""Burndown: a sprint's remaining story points over time, from Jira's scope change data."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jira_exporter.sprint_report import tidy_points

COLUMNS = ["Date", "Remaining Story Points", "Guideline", "Completed That Day",
           "Scope Change That Day"]


class BurndownError(Exception):
    pass


@dataclass
class BurndownDay:
    day: date
    remaining: float
    guideline: float
    completed: float
    scope_change: float

    def values(self) -> list:
        return [self.day, tidy_points(self.remaining), tidy_points(round(self.guideline, 1)),
                tidy_points(self.completed), tidy_points(self.scope_change)]


@dataclass
class Burndown:
    start: datetime  # local time
    end: datetime  # planned end
    stop: datetime  # when the chart stops: completion, or now for an active sprint
    completed: bool  # the sprint was closed
    start_remaining: float
    # Remaining story points after each change from the start to `stop`, as a step line.
    steps: list[tuple[datetime, float]] = field(default_factory=list)
    days: list[BurndownDay] = field(default_factory=list)

    @property
    def remaining(self) -> float:
        return self.steps[-1][1] if self.steps else self.start_remaining


def _when(ms) -> datetime:
    """Jira's milliseconds since 1970 (UTC) as local time without a time zone."""
    return datetime.fromtimestamp(int(ms) / 1000).replace(microsecond=0)


def _float(value) -> float:
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


class _Sprint:
    """The sprint's issues as the changes are replayed: in the sprint, estimate, done."""

    def __init__(self):
        self.issues: dict[str, dict] = {}

    def apply(self, change: dict):
        key = change.get("key")
        if not key:
            return
        issue = self.issues.setdefault(key, {"in": False, "est": 0.0, "done": False})
        if "added" in change:
            issue["in"] = bool(change["added"])
        stat = change.get("statC")
        if isinstance(stat, dict) and ("newValue" in stat or "oldValue" in stat):
            issue["est"] = _float(stat.get("newValue"))
        column = change.get("column")
        if isinstance(column, dict):
            if column.get("done"):
                issue["done"] = True
            elif column.get("notDone"):
                issue["done"] = False

    def remaining(self) -> float:
        return sum(i["est"] for i in self.issues.values() if i["in"] and not i["done"])

    def done(self) -> float:
        return sum(i["est"] for i in self.issues.values() if i["in"] and i["done"])

    def scope(self) -> float:
        return sum(i["est"] for i in self.issues.values() if i["in"])


def build_burndown(data: dict) -> Burndown:
    """Replay Jira's scope change data: changes up to the sprint's start set the committed
    work, later ones move the remaining line. Days run from the start to the sprint's
    completion (or Jira's 'now' while it is still active)."""
    try:
        start = _when(data["startTime"])
        end = _when(data["endTime"])
    except (KeyError, TypeError, ValueError):
        raise BurndownError("Jira didn't return a start and end for this sprint; it may "
                            "not have started yet.") from None
    completed = bool(data.get("completeTime"))
    if completed:
        stop = _when(data["completeTime"])
    elif data.get("now"):
        stop = _when(data["now"])
    else:
        stop = end
    stop = max(stop, start)

    changes = []
    for ms, entries in (data.get("changes") or {}).items():
        try:
            when = _when(ms)
        except (TypeError, ValueError):
            continue
        for entry in entries if isinstance(entries, list) else []:
            if isinstance(entry, dict):
                changes.append((when, entry))
    changes.sort(key=lambda c: c[0])  # stable: same-time changes keep Jira's order

    sprint = _Sprint()
    i = 0
    while i < len(changes) and changes[i][0] <= start:
        sprint.apply(changes[i][1])
        i += 1
    chart = Burndown(start, end, stop, completed, sprint.remaining())
    chart.steps.append((start, chart.start_remaining))

    day = start.date()
    while day <= stop.date():
        day_end = min(datetime.combine(day + timedelta(days=1), time()), stop)
        done_before, scope_before = sprint.done(), sprint.scope()
        while i < len(changes) and changes[i][0] <= day_end:
            when, entry = changes[i]
            sprint.apply(entry)
            if sprint.remaining() != chart.steps[-1][1]:
                chart.steps.append((when, sprint.remaining()))
            i += 1
        chart.days.append(BurndownDay(
            day, sprint.remaining(), _guideline(chart, day_end),
            sprint.done() - done_before, sprint.scope() - scope_before))
        day += timedelta(days=1)
    return chart


def _guideline(chart: Burndown, when: datetime) -> float:
    """Straight line from the committed work at the start to zero at the planned end."""
    span = (chart.end - chart.start).total_seconds()
    if span <= 0:
        return 0.0
    left = 1 - (when - chart.start).total_seconds() / span
    return max(0.0, chart.start_remaining * left)


def write_burndown_xlsx(path: Path, title: str, chart: Burndown) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Burndown"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for day in chart.days:
        ws.append(day.values())
    last = ws.max_row
    for r in range(2, last + 1):
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
    for i, header in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(header), 10) + 2
    ws.freeze_panes = "A2"
    if chart.days:
        line = LineChart()
        line.title = f"{title} burndown"
        line.y_axis.title = "Story points"
        line.x_axis.title = "Date"
        line.x_axis.number_format = "mm-dd"
        line.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=last),
                      titles_from_data=True)
        line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        line.width, line.height = 24, 12
        ws.add_chart(line, f"{get_column_letter(len(COLUMNS) + 2)}2")
    wb.save(path)
