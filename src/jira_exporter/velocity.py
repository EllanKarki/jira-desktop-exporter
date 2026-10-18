"""Velocity: committed vs completed story points per sprint for one team, as an Excel chart."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from jira_exporter.sprint_report import SprintRow, tidy_points

COLUMNS = ["Sprint Name", "Committed Story Points", "Completed Story Points",
           "% Completed vs Committed"]


def average_velocity(rows: list[SprintRow]) -> float | None:
    """Mean completed story points per sprint; None with no sprints."""
    if not rows:
        return None
    return sum(r.completed_points for r in rows) / len(rows)


def velocity_values(row: SprintRow) -> list:
    return [row.sprint, tidy_points(row.committed_points), tidy_points(row.completed_points),
            row.percent]


def write_velocity_xlsx(path: Path, team: str, rows: list[SprintRow]) -> None:
    """A table of the sprints with a clustered bar chart of committed vs completed."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Velocity"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append(velocity_values(row))
    last = ws.max_row
    for r in range(2, last + 1):
        ws.cell(r, len(COLUMNS)).number_format = "0%"
    average = average_velocity(rows)
    ws.append([])
    ws.append(["Average velocity", None,
               None if average is None else round(average, 2)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for i, header in enumerate(COLUMNS, start=1):
        width = max([len(header)] + [len(str(velocity_values(r)[i - 1])) for r in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60) + 2
    ws.freeze_panes = "A2"

    if rows:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.title = f"{team} velocity"
        chart.y_axis.title = "Story points"
        chart.x_axis.title = "Sprint"
        data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        chart.width, chart.height = 24, 12
        ws.add_chart(chart, f"{get_column_letter(len(COLUMNS) + 2)}2")
    wb.save(path)
