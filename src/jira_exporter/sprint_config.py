"""Team sprint settings from an Excel file, and the sprint they produce for a given number."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "Teams"
# Column headers, in template order. Matching ignores case and extra spaces.
TEAM = "Team"
PROJECT = "Project Key"
BOARD = "Board ID"
NAME_FORMAT = "Sprint Name Format"
FIRST_NUMBER = "First Sprint Number"
FIRST_START = "First Sprint Start Date"
LENGTH = "Sprint Length (days)"
START_TIME = "Start Time"
GOAL = "Sprint Goal"
COLUMNS = [TEAM, PROJECT, BOARD, NAME_FORMAT, FIRST_NUMBER, FIRST_START, LENGTH,
           START_TIME, GOAL]
REQUIRED = [TEAM, PROJECT, BOARD, FIRST_NUMBER, FIRST_START, LENGTH]

DEFAULT_NAME_FORMAT = "{team} Sprint {number}"
DEFAULT_START_TIME = time(9, 0)


class SprintConfigError(Exception):
    pass


@dataclass
class TeamConfig:
    team: str
    project_key: str
    board_id: int
    first_number: int
    first_start: date
    length_days: int
    name_format: str = DEFAULT_NAME_FORMAT
    start_time: time = DEFAULT_START_TIME
    goal: str = ""
    row: int = 0  # Excel row, for messages


@dataclass
class SprintPlan:
    name: str
    board_id: int
    start: datetime  # local time
    end: datetime
    goal: str


def plan_sprint(cfg: TeamConfig, number: int) -> SprintPlan:
    """Sprint `number` starts (number - first) sprint lengths after the first sprint starts,
    and ends when the next one starts."""
    if number < 1:
        raise SprintConfigError("The sprint number must be 1 or more.")
    start = datetime.combine(cfg.first_start, cfg.start_time) + \
        timedelta(days=(number - cfg.first_number) * cfg.length_days)
    end = start + timedelta(days=cfg.length_days)
    values = {"team": cfg.team, "project": cfg.project_key, "number": number}
    return SprintPlan(_fill(cfg.name_format, values, cfg.row, NAME_FORMAT).strip(),
                      cfg.board_id, start, end,
                      _fill(cfg.goal, values, cfg.row, GOAL).strip())


def jira_datetime(dt: datetime) -> str:
    """ISO 8601 with the computer's UTC offset, e.g. 2026-01-05T09:00:00.000+05:30."""
    if dt.tzinfo is None:
        dt = dt.astimezone()  # treat as local time
    return dt.isoformat(timespec="milliseconds")


def _fill(template: str, values: dict, row: int, column: str) -> str:
    try:
        return template.format(**values)
    except (KeyError, IndexError, ValueError) as exc:
        raise SprintConfigError(
            f"Row {row}, '{column}': can't use '{template}' ({exc}). Use only "
            "{team}, {project} and {number} as placeholders.") from exc


def load_teams(path: Path) -> tuple[list[TeamConfig], list[str]]:
    """Read the Teams sheet. Returns the valid rows and a message for each skipped row."""
    path = Path(path)
    if not path.is_file():
        raise SprintConfigError(f"Sprint config not found: {path}")
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # openpyxl raises several types for non-Excel files
        raise SprintConfigError(f"{path} is not a readable .xlsx file: {exc}") from exc
    try:
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise SprintConfigError(f"{path}: the '{ws.title}' sheet is empty.")

    header = {_norm(h): i for i, h in enumerate(rows[0]) if h is not None}
    missing = [c for c in REQUIRED if _norm(c) not in header]
    if missing:
        raise SprintConfigError(f"{path}: missing column(s): {', '.join(missing)}. "
                                f"Expected: {', '.join(COLUMNS)}.")

    def cell(values, column):
        i = header.get(_norm(column))
        return values[i] if i is not None and i < len(values) else None

    teams, problems = [], []
    for n, values in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        try:
            teams.append(_parse_row(n, lambda c: cell(values, c)))
        except SprintConfigError as exc:
            problems.append(str(exc))
    unique, seen = [], set()
    for t in teams:
        key = (t.team.lower(), t.project_key.lower())
        if key in seen:
            problems.append(f"Row {t.row}: {t.team} / {t.project_key} is listed more than "
                            "once; the first row is used.")
            continue
        seen.add(key)
        unique.append(t)
    return unique, problems


def _parse_row(n: int, get) -> TeamConfig:
    def text(column):
        v = get(column)
        return "" if v is None else str(v).strip()

    for column in REQUIRED:
        if text(column) == "":
            raise SprintConfigError(f"Row {n}: '{column}' is empty.")
    board = _int(get(BOARD), n, BOARD, minimum=1)
    first = _int(get(FIRST_NUMBER), n, FIRST_NUMBER, minimum=1)
    length = _int(get(LENGTH), n, LENGTH, minimum=1)
    return TeamConfig(
        team=text(TEAM),
        project_key=text(PROJECT),
        board_id=board,
        first_number=first,
        first_start=_date(get(FIRST_START), n),
        length_days=length,
        name_format=text(NAME_FORMAT) or DEFAULT_NAME_FORMAT,
        start_time=_time(get(START_TIME), n),
        goal=text(GOAL),
        row=n,
    )


def _norm(header) -> str:
    return " ".join(str(header).split()).lower()


def _int(value, n, column, minimum) -> int:
    try:
        number = float(str(value).strip())
        if number != int(number):
            raise ValueError
        number = int(number)
    except (TypeError, ValueError):
        raise SprintConfigError(f"Row {n}: '{column}' must be a whole number, not '{value}'.")
    if number < minimum:
        raise SprintConfigError(f"Row {n}: '{column}' must be {minimum} or more.")
    return number


def _date(value, n) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    raise SprintConfigError(f"Row {n}: '{FIRST_START}' must be a date (for example "
                            f"2026-01-05), not '{value}'.")


def _time(value, n) -> time:
    if value is None or str(value).strip() == "":
        return DEFAULT_START_TIME
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)) and 0 <= value < 1:  # Excel stores times as day fractions
        minutes = round(value * 24 * 60)
        return time(minutes // 60 % 24, minutes % 60)
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p"):
        try:
            return datetime.strptime(str(value).strip().upper(), fmt).time()
        except ValueError:
            pass
    raise SprintConfigError(f"Row {n}: '{START_TIME}' must be a time like 09:00, not '{value}'.")


EXAMPLE_ROWS = [
    ["Team Alpha", "ABC", 1, "{team} Sprint {number}", 1, date(2026, 1, 5), 14, time(9, 0),
     "Sprint {number} goal"],
    ["Team Beta", "ABC", 2, "ABC Beta {number}", 1, date(2026, 1, 12), 7, time(10, 0), ""],
    ["Team Gamma", "XYZ", 3, "", 20, date(2026, 3, 2), 21, None, ""],
]
NOTES = [
    ("Team", "Name shown in the Create Sprint form. A team can have one row per project."),
    ("Project Key", "Jira project key, e.g. ABC."),
    ("Board ID", "The Scrum board the sprint is created on. Open the board in Jira: the "
                 "number after rapidView= or /boards/ in the address is its ID."),
    ("Sprint Name Format", "Optional. Placeholders: {team}, {project}, {number}. "
                           f"Default: {DEFAULT_NAME_FORMAT}"),
    ("First Sprint Number", "The number of a sprint whose start date you know."),
    ("First Sprint Start Date", "That sprint's start date. Later sprints follow on back to "
                                "back, one sprint length apart."),
    ("Sprint Length (days)", "e.g. 14 for two-week sprints. A sprint ends when the next starts."),
    ("Start Time", "Optional, e.g. 09:00 (your computer's time zone). Default 09:00."),
    ("Sprint Goal", "Optional. Same placeholders as the name."),
]


def write_template(path: Path) -> None:
    """Write an example sprint config with the expected columns and a Help sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for i, width in enumerate([16, 12, 10, 26, 18, 22, 20, 11, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for r in range(2, ws.max_row + 1):
        ws.cell(r, COLUMNS.index(FIRST_START) + 1).number_format = "yyyy-mm-dd"
        ws.cell(r, COLUMNS.index(START_TIME) + 1).number_format = "hh:mm"
    ws.freeze_panes = "A2"
    help_ws = wb.create_sheet("Help")
    help_ws.append(["Column", "What to enter"])
    for cell in help_ws[1]:
        cell.font = Font(bold=True)
    for row in NOTES:
        help_ws.append(list(row))
    help_ws.column_dimensions["A"].width = 24
    help_ws.column_dimensions["B"].width = 100
    wb.save(path)
