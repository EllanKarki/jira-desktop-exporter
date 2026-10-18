"""Sprint report: per team and sprint, what was committed, completed, spilled and changed."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    "Team Name",
    "Sprint Name",
    "Committed Issues",
    "Committed Story Points",
    "Completed Issues",
    "Completed Story Points",
    "Spill Over Issues",
    "Spill Over Story Points",
    "Issues Adjusted During Sprint",
    "% Completed vs Committed",
]


@dataclass
class SprintRow:
    team: str
    sprint: str
    committed: int
    committed_points: float
    completed: int
    completed_points: float
    spilled: int
    spilled_points: float
    added: int
    removed: int
    board_id: int = 0
    sprint_id: int = 0

    @property
    def adjusted(self) -> int:
        """Issues added after the sprint started plus issues removed from it."""
        return self.added + self.removed

    @property
    def percent(self) -> float | None:
        """Completed story points as a fraction of committed ones (0.85 = 85%); None when
        nothing was committed. Work added mid-sprint can take it over 100%."""
        if not self.committed_points:
            return None
        return self.completed_points / self.committed_points

    def values(self) -> list:
        """The row in COLUMNS order."""
        return [self.team, self.sprint, self.committed, tidy_points(self.committed_points),
                self.completed, tidy_points(self.completed_points), self.spilled,
                tidy_points(self.spilled_points), self.adjusted, self.percent]


def tidy_points(points: float):
    """3.0 -> 3, so whole story points don't show a trailing .0."""
    return int(points) if float(points).is_integer() else round(points, 2)


def _estimate(issue: dict, key: str) -> float:
    stat = issue.get(key) or {}
    value = (stat.get("statFieldValue") or {}).get("value")
    try:
        return float(value) if value is not None else 0.0
    except (TypeError, ValueError):
        return 0.0


def estimate_field(report: dict) -> str | None:
    """The field id the board estimates with (e.g. customfield_10002), if Jira says."""
    contents = report.get("contents") or {}
    for key in ("completedIssues", "issuesNotCompletedInCurrentSprint", "puntedIssues"):
        for issue in contents.get(key) or []:
            field = (issue.get("estimateStatistic") or {}).get("statFieldId")
            if field:
                return field
    return None


def summarize(team: str, sprint: dict, report: dict, board_id: int = 0) -> SprintRow:
    """Turn Jira's sprint report into one row.

    - Committed: issues in the sprint when it started, with their estimate at the start.
    - Completed: issues done when the sprint closed, with their final estimate. This
      includes issues added after the start, like Jira's own report.
    - Spill over: issues still not done when the sprint closed.
    - Adjusted: issues added after the start plus issues removed before the end.
    Issues Jira lists as "completed outside this sprint" are left out of every column.
    """
    contents = report.get("contents") or {}
    completed = contents.get("completedIssues") or []
    not_done = contents.get("issuesNotCompletedInCurrentSprint") or []
    removed = contents.get("puntedIssues") or []
    added_keys = contents.get("issueKeysAddedDuringSprint") or {}
    added = set(added_keys) if isinstance(added_keys, (dict, list)) else set()

    committed = [i for i in completed + not_done + removed if i.get("key") not in added]
    return SprintRow(
        team=team,
        sprint=sprint.get("name") or (report.get("sprint") or {}).get("name", ""),
        committed=len(committed),
        committed_points=sum(_estimate(i, "estimateStatistic") for i in committed),
        completed=len(completed),
        completed_points=sum(_current(i) for i in completed),
        spilled=len(not_done),
        spilled_points=sum(_current(i) for i in not_done),
        added=len(added),
        removed=len(removed),
        board_id=board_id,
        sprint_id=int(sprint.get("id") or 0),
    )


def _current(issue: dict) -> float:
    if (issue.get("currentEstimateStatistic") or {}).get("statFieldValue") is not None:
        return _estimate(issue, "currentEstimateStatistic")
    return _estimate(issue, "estimateStatistic")


def recent_closed_sprints(sprints: list[dict], board_id: int, count: int) -> list[dict]:
    """The last `count` closed sprints that belong to the board, oldest first.

    A board's sprint list also shows sprints from other boards that share its issues;
    those are left for their own board."""
    own = [s for s in sprints if (s.get("state") or "").lower() == "closed"
           and s.get("originBoardId") in (None, board_id)]
    own.sort(key=lambda s: (s.get("completeDate") or s.get("endDate") or "", s.get("id") or 0))
    return own[-count:] if count > 0 else []


def write_report_xlsx(path: Path, rows: list[SprintRow]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint Report"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        ws.append(row.values())
    percent_col = len(COLUMNS)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, percent_col).number_format = "0%"
    for i, header in enumerate(COLUMNS, start=1):
        width = max([len(header)] + [len(str(row.values()[i - 1])) for row in rows])
        ws.column_dimensions[get_column_letter(i)].width = min(width, 60) + 2
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
